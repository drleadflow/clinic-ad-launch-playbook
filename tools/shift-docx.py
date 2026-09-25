#!/usr/bin/env python3
"""Markdown -> styled Word doc (US Letter, Calibri, navy headings 1F3A5F, red subheads C8102E,
shaded script boxes with a navy left border, tables with a navy header row). Built for the Shift
Framework file but works for any campaign markdown.

usage: shift-docx.py <in.md> <out.docx> [--title "Doc title"] [--subtitle "..."] [--prepared "Prepared by ..."]

Markdown handled: # / ## / ### headings, paragraphs, - bullets, 1. numbered, | tables |,
> blockquotes (rendered as a shaded script box; consecutive > lines = one box; a first line in
**bold** becomes the box label), --- rules, **bold** and *italic* inline. Em dashes are refused
(exit 2) so the writing rule is enforced at export.
"""
import re, sys
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

NAVY, RED, SHADE = "1F3A5F", "C8102E", "F2F4F7"

def arg(flag, default=None):
    if flag in sys.argv:
        i = sys.argv.index(flag); return sys.argv[i + 1]
    return default

src, out = sys.argv[1], sys.argv[2]
text = open(src, encoding="utf-8").read()
if "—" in text:
    print("refused: em dash found. Replace with a period, comma or colon.", file=sys.stderr); sys.exit(2)

d = Document()
for s in d.sections:
    s.page_width, s.page_height = Inches(8.5), Inches(11)
    s.left_margin = s.right_margin = s.top_margin = s.bottom_margin = Inches(1)
st = d.styles["Normal"]; st.font.name = "Calibri"; st.font.size = Pt(11)
st.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")

def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr(); e = OxmlElement("w:shd")
    e.set(qn("w:val"), "clear"); e.set(qn("w:color"), "auto"); e.set(qn("w:fill"), fill); tcPr.append(e)

def inline(par, s, size=None, color=None, bold=None, italic=None):
    """**bold** and *italic* runs."""
    for tok in re.split(r"(\*\*[^*]+\*\*|\*[^*]+\*)", s):
        if not tok: continue
        b, i = bold, italic; t = tok
        if tok.startswith("**"): t, b = tok[2:-2], True
        elif tok.startswith("*"): t, i = tok[1:-1], True
        r = par.add_run(t)
        if b: r.bold = True
        if i: r.italic = True
        if size: r.font.size = Pt(size)
        if color: r.font.color.rgb = RGBColor.from_string(color)

def heading(t, lvl):
    p = d.add_heading(level=min(lvl, 3)); p.text = ""
    r = p.add_run(t); r.font.name = "Calibri"; r.bold = True
    r.font.size = Pt({1: 18, 2: 13, 3: 12}[min(lvl, 3)])
    r.font.color.rgb = RGBColor.from_string(NAVY if lvl == 1 else RED)

def box(lines):
    t = d.add_table(rows=1, cols=1); c = t.rows[0].cells[0]; shade(c, SHADE)
    tcPr = c._tc.get_or_add_tcPr(); b = OxmlElement("w:tcBorders")
    for side in ("left", "top", "bottom", "right"):
        e = OxmlElement(f"w:{side}")
        if side == "left": e.set(qn("w:val"), "single"); e.set(qn("w:sz"), "24"); e.set(qn("w:color"), NAVY)
        else: e.set(qn("w:val"), "nil")
        b.append(e)
    tcPr.append(b)
    first = True
    for ln in lines:
        p = c.paragraphs[0] if first else c.add_paragraph(); first = False
        p.paragraph_format.space_after = Pt(0)
        m = re.fullmatch(r"\*\*(.+)\*\*", ln.strip())
        if m: r = p.add_run(m.group(1)); r.bold = True; r.font.color.rgb = RGBColor.from_string(NAVY)
        else: inline(p, ln)
    d.add_paragraph()

def table(rows):
    t = d.add_table(rows=1, cols=len(rows[0])); t.style = "Table Grid"
    for i, h in enumerate(rows[0]):
        c = t.rows[0].cells[i]; shade(c, NAVY); c.paragraphs[0].text = ""
        inline(c.paragraphs[0], h, size=10, color="FFFFFF", bold=True)
    for row in rows[1:]:
        cs = t.add_row().cells
        for i, v in enumerate(row[:len(rows[0])]):
            cs[i].paragraphs[0].text = ""; inline(cs[i].paragraphs[0], v, size=10)
    d.add_paragraph()

# optional title block
if arg("--title"):
    p = d.add_paragraph(); r = p.add_run(arg("--title")); r.font.size = Pt(24); r.bold = True; r.font.color.rgb = RGBColor.from_string(NAVY)
    if arg("--subtitle"):
        p = d.add_paragraph(); r = p.add_run(arg("--subtitle")); r.font.size = Pt(15); r.font.color.rgb = RGBColor.from_string(RED)
    if arg("--prepared"): d.add_paragraph(arg("--prepared"))

lines = text.splitlines(); i = 0
while i < len(lines):
    ln = lines[i]
    if not ln.strip(): i += 1; continue
    if ln.startswith("---"): i += 1; continue
    m = re.match(r"^(#{1,3})\s+(.*)", ln)
    if m: heading(m.group(2).strip(), len(m.group(1))); i += 1; continue
    if ln.lstrip().startswith(">"):
        blk = []
        while i < len(lines) and lines[i].lstrip().startswith(">"):
            blk.append(re.sub(r"^\s*>\s?", "", lines[i])); i += 1
        box([b for b in blk if b.strip() != ""]); continue
    if ln.lstrip().startswith("|"):
        rows = []
        while i < len(lines) and lines[i].lstrip().startswith("|"):
            cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
            if not all(re.fullmatch(r":?-+:?", c) for c in cells if c): rows.append(cells)
            i += 1
        if rows: table(rows)
        continue
    m = re.match(r"^\s*[-*]\s+(.*)", ln)
    if m: p = d.add_paragraph(style="List Bullet"); inline(p, m.group(1)); i += 1; continue
    m = re.match(r"^\s*\d+[.)]\s+(.*)", ln)
    if m: p = d.add_paragraph(style="List Number"); inline(p, m.group(1)); i += 1; continue
    p = d.add_paragraph(); inline(p, ln.strip()); i += 1

d.save(out)
n = sum(1 for _ in d.paragraphs); print(f"ok {out}: {n} paragraphs, {len(d.tables)} tables, 0 em dashes")
